from pathlib import Path


def load_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name == "Heading 1":
            parts.append(f"# {text}")
        elif style_name == "Heading 2":
            parts.append(f"## {text}")
        elif style_name == "Heading 3":
            parts.append(f"### {text}")
        else:
            parts.append(text)

    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [_escape_cell(cell.text.strip()) for cell in row.cells]
            if any(cells):
                rows.append(f"| {' | '.join(cells)} |")
        if rows:
            parts.append("\n".join(rows))

    return "\n\n".join(parts)


def load_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    parts: list[str] = []

    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = [_format_cell(value) for value in row]
            if any(values):
                rows.append(f"| {' | '.join(_escape_cell(value) for value in values)} |")
        if rows:
            parts.append(f"## {sheet.title}\n\n" + "\n".join(rows))

    workbook.close()
    return "\n\n".join(parts)


def _format_cell(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").strip()


def _escape_cell(value: str) -> str:
    return value.replace("|", "\\|")
